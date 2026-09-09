"""
Flask application for automated Google Sheets Academic Journal creation.
Phase 1: File upload, parsing, and preview (Excel, CSV, Word, Google Doc).
Phase 2: Google Sheets integration — create and share journals.
"""

import csv
import datetime
import json
import logging
import os
import re
import tempfile
import uuid
from pathlib import Path

import pandas as pd
from flask import Flask, flash, redirect, render_template, request, session, url_for
from werkzeug.utils import secure_filename

from google_sheets import create_academic_journal, get_drive_credentials, get_service_account_email
from doc_import import match_groups_to_rnp, parse_rnp, parse_student_roster

logger = logging.getLogger(__name__)

app = Flask(__name__)
app.secret_key = os.urandom(32)

ALLOWED_EXTENSIONS = {"xlsx", "xls", "csv", "docx"}
ALLOWED_ROSTER_RNP_EXTENSIONS = {"doc", "docx"}

DATA_DIR = Path(__file__).parent / "data"
ROSTER_PATH = DATA_DIR / "roster.json"
BATCH_STATE_DIR = DATA_DIR / "batch_state"


def allowed_file(filename: str) -> bool:
    """Check if the uploaded file has an allowed extension."""
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def extract_google_doc_id(url: str) -> str | None:
    """Extract Document ID from a Google Docs URL."""
    match = re.search(r"/document/d/([a-zA-Z0-9-_]+)", url)
    if match:
        return match.group(1)
    return None


def parse_docx_file(filepath: str) -> dict:
    """
    Parse an uploaded Word (.docx) file containing group list table.
    Expects a table containing student names (under a PІБ column) and disciplines.
    """
    try:
        import docx
    except ImportError:
        return {"error": "Бібліотеку 'python-docx' не встановлено на сервері."}

    try:
        doc = docx.Document(filepath)
        
        # 1. Search for Group Name in paragraphs and tables
        group_name = None
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            text_lower = text.lower()
            if any(kw in text_lower for kw in ["група", "group", "назва"]):
                if ":" in text:
                    group_name = text.split(":", 1)[1].strip()
                    break
                else:
                    words = text.split()
                    for idx, w in enumerate(words):
                        if any(kw in w.lower() for kw in ["група", "group"]):
                            if idx + 1 < len(words):
                                group_name = words[idx + 1].strip(" -:")
                                break
                    if group_name:
                        break

        # 2. Locate students/disciplines table
        students_table = None
        student_col_idx = None
        header_row_idx = 0
        student_col_keywords = ["піб", "прізвище", "студент", "ім'я", "name"]

        for table in doc.tables:
            for r_idx in range(min(len(table.rows), 3)):
                row_cells = [cell.text.strip().lower() for cell in table.rows[r_idx].cells]
                for idx, cell_text in enumerate(row_cells):
                    if any(kw in cell_text for kw in student_col_keywords):
                        students_table = table
                        student_col_idx = idx
                        header_row_idx = r_idx
                        break
                if students_table:
                    break
            if students_table:
                break

        if not students_table:
            return {"error": "Не знайдено таблицю зі списком студентів та дисциплін у документі Word."}

        # Retrieve group name from table if not found in paragraphs
        if not group_name:
            for r_idx in range(header_row_idx):
                cells = [c.text.strip() for c in students_table.rows[r_idx].cells]
                for idx, val in enumerate(cells):
                    if any(kw in val.lower() for kw in ["група", "group"]):
                        if idx + 1 < len(cells):
                            group_name = cells[idx + 1].strip(" -:")
                            break
                        
        if not group_name:
            group_name = "Не визначено"

        # 4. Parse headers
        header_cells = [c.text.strip() for c in students_table.rows[header_row_idx].cells]
        
        # 5. Extract Students and Disciplines
        students = []
        discipline_cols = []
        skip_keywords = ["№", "#", "номер", "num", header_cells[student_col_idx].lower()]

        for idx, col_name in enumerate(header_cells):
            if col_name.lower().strip() not in skip_keywords and idx != student_col_idx:
                discipline_cols.append((idx, col_name))

        disciplines_data = {col_idx: [] for col_idx, _ in discipline_cols}

        # Parse student rows
        for r_idx in range(header_row_idx + 1, len(students_table.rows)):
            row = students_table.rows[r_idx]
            cells = [c.text.strip() for c in row.cells]
            
            if len(cells) <= student_col_idx or not cells[student_col_idx]:
                continue
                
            students.append(cells[student_col_idx])
            
            for col_idx, _ in discipline_cols:
                val = cells[col_idx] if col_idx < len(cells) else ""
                disciplines_data[col_idx].append(val)

        # Build disciplines metadata
        disciplines = []
        for col_idx, col_name in discipline_cols:
            count = None
            for val in disciplines_data[col_idx]:
                if val.isdigit():
                    count = int(val)
                    break
            
            col_lower = col_name.lower().strip()
            is_cp = False
            for kw in ["курсов", "course project", "coursework", "course work"]:
                if kw in col_lower:
                    is_cp = True
                    break
            if not is_cp:
                tokens = col_lower.replace("(", " ").replace(")", " ").replace(".", " ").replace(",", " ").split()
                if any(t in tokens for t in ["кп", "кр", "kp", "kr"]):
                    is_cp = True

            if is_cp:
                control_type = "course_project"
            else:
                is_exam = any(kw in col_lower for kw in ["екзамен", "екз", "exam"])
                control_type = "exam" if is_exam else "credit"

            disciplines.append({
                "name": col_name,
                "class_count": count,
                "control_type": control_type
            })

        if not students:
            return {"error": "У таблиці документа Word не знайдено жодного студента."}

        return {
            "group_name": group_name,
            "students": students,
            "disciplines": disciplines,
            "error": None
        }
    except Exception as exc:
        return {"error": f"Помилка при зчитуванні Word-файлу: {exc}"}


def _read_csv_ragged(filepath: str, header_row: int | None) -> pd.DataFrame:
    """Read a CSV into a DataFrame, tolerating rows with a different number of
    columns than the rest (pandas' C parser raises on this — real exports often
    have a short "Назва групи" row above a wider student table)."""
    with open(filepath, "r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.reader(f))

    max_cols = max((len(row) for row in rows), default=0)
    padded = [row + [""] * (max_cols - len(row)) for row in rows]

    if header_row is None:
        df = pd.DataFrame(padded, dtype=str)
    else:
        columns = padded[header_row] if header_row < len(padded) else list(range(max_cols))
        df = pd.DataFrame(padded[header_row + 1:], columns=columns, dtype=str)

    return df.replace("", pd.NA)


def parse_uploaded_file(filepath: str) -> dict:
    """Parse an uploaded Excel, CSV, or Word file."""
    ext = Path(filepath).suffix.lower()

    if ext == ".docx":
        return parse_docx_file(filepath)

    try:
        if ext == ".csv":
            raw_df = _read_csv_ragged(filepath, header_row=None)
        else:
            raw_df = pd.read_excel(filepath, header=None, dtype=str, engine="openpyxl")

        if raw_df.empty:
            return {"error": "Файл порожній. Завантажте файл із даними."}

        group_name = None
        header_row_idx = 0

        first_row_values = raw_df.iloc[0].dropna().tolist()
        for col_idx in range(len(first_row_values)):
            cell = str(first_row_values[col_idx]).strip().lower()
            if any(kw in cell for kw in ["група", "group", "назва"]):
                if col_idx + 1 < len(first_row_values):
                    group_name = str(first_row_values[col_idx + 1]).strip()
                break

        if group_name is None and len(first_row_values) == 2:
            group_name = str(first_row_values[1]).strip()

        if group_name is None:
            group_name = "Не визначено"

        student_col_keywords = ["піб", "прізвище", "студент", "ім'я", "name"]
        header_found = False

        for row_idx in range(raw_df.shape[0]):
            row_cells = raw_df.iloc[row_idx].dropna().astype(str).str.strip().str.lower()
            for cell in row_cells:
                if any(kw in cell for kw in student_col_keywords):
                    header_row_idx = row_idx
                    header_found = True
                    break
            if header_found:
                break

        if not header_found:
            header_row_idx = 1 if raw_df.shape[0] > 1 else 0

        if ext == ".csv":
            df = _read_csv_ragged(filepath, header_row=header_row_idx)
        else:
            df = pd.read_excel(filepath, header=header_row_idx, dtype=str, engine="openpyxl")

        df.columns = df.columns.astype(str).str.strip()

        student_col = None
        for col in df.columns:
            if any(kw in col.lower() for kw in student_col_keywords):
                student_col = col
                break

        if student_col is None:
            student_col = df.columns[1] if len(df.columns) > 1 else df.columns[0]

        students = (
            df[student_col]
            .dropna()
            .astype(str)
            .str.strip()
            .loc[lambda s: s != ""]
            .tolist()
        )

        skip_keywords = ["№", "#", "номер", "num", student_col.lower()]
        discipline_cols = [
            col
            for col in df.columns
            if col.lower().strip() not in skip_keywords and col != student_col
        ]

        disciplines: list[dict] = []
        for col in discipline_cols:
            count = None
            series = pd.to_numeric(df[col], errors="coerce").dropna()
            if not series.empty:
                count = int(series.iloc[0])

            col_lower = col.lower().strip()
            is_cp = False
            for kw in ["курсов", "course project", "coursework", "course work"]:
                if kw in col_lower:
                    is_cp = True
                    break
            if not is_cp:
                tokens = col_lower.replace("(", " ").replace(")", " ").replace(".", " ").replace(",", " ").split()
                if any(t in tokens for t in ["кп", "кр", "kp", "kr"]):
                    is_cp = True

            if is_cp:
                control_type = "course_project"
            else:
                is_exam = any(kw in col_lower for kw in ["екзамен", "екз", "exam"])
                control_type = "exam" if is_exam else "credit"

            disciplines.append({
                "name": col,
                "class_count": count,
                "control_type": control_type
            })

        return {
            "group_name": group_name,
            "students": students,
            "disciplines": disciplines,
            "error": None,
        }

    except Exception as exc:
        return {"error": f"Помилка під час обробки файлу: {exc}"}


# ==========================================================================
# Roster & batch-state persistence (Список студентів + РНП flow)
# ==========================================================================

def load_roster() -> dict:
    """Load the persisted student roster ({group_name: [student, ...]})."""
    if not ROSTER_PATH.exists():
        return {}
    try:
        with open(ROSTER_PATH, "r", encoding="utf-8") as f:
            return json.load(f).get("groups", {})
    except Exception:
        return {}


def save_roster(groups: dict) -> None:
    """Persist the student roster to disk so it survives across sessions/restarts."""
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    with open(ROSTER_PATH, "w", encoding="utf-8") as f:
        json.dump(
            {"updated_at": datetime.datetime.now().isoformat(), "groups": groups},
            f, ensure_ascii=False, indent=2,
        )


def save_batch_state(batch_id: str, match: dict) -> None:
    """Persist a matched-groups batch to disk (kept out of the session cookie, which
    is too small for many groups/students/disciplines)."""
    BATCH_STATE_DIR.mkdir(parents=True, exist_ok=True)
    with open(BATCH_STATE_DIR / f"{batch_id}.json", "w", encoding="utf-8") as f:
        json.dump(match, f, ensure_ascii=False)


def load_batch_state(batch_id: str) -> dict | None:
    path = BATCH_STATE_DIR / f"{batch_id}.json"
    if not path.exists():
        return None
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return None


def delete_batch_state(batch_id: str) -> None:
    path = BATCH_STATE_DIR / f"{batch_id}.json"
    try:
        path.unlink(missing_ok=True)
    except OSError:
        pass


def _save_upload_to_tmp(file_storage) -> str:
    """Save a Flask FileStorage to a fresh temp dir and return the filepath."""
    tmp_dir = tempfile.mkdtemp()
    filepath = os.path.join(tmp_dir, secure_filename(file_storage.filename))
    file_storage.save(filepath)
    return filepath


def _cleanup_tmp(filepath: str) -> None:
    try:
        tmp_dir = os.path.dirname(filepath)
        os.remove(filepath)
        os.rmdir(tmp_dir)
    except OSError:
        pass


# ==========================================================================
# Routes
# ==========================================================================

@app.route("/", methods=["GET"])
def upload_page():
    """Render the file-upload page, displaying the service email if configured."""
    service_email = get_service_account_email() or ""

    roster = load_roster()
    roster_info = None
    if roster:
        roster_info = {
            "group_count": len(roster),
            "student_count": sum(len(v) for v in roster.values()),
            "groups": sorted(roster.keys()),
        }

    return render_template("upload.html", service_email=service_email, roster_info=roster_info)


@app.route("/upload", methods=["POST"])
def handle_upload():
    """Receive the uploaded file, parse it, and redirect to preview."""
    if "file" not in request.files:
        flash("Файл не було надіслано.", "error")
        return redirect(url_for("upload_page"))

    file = request.files["file"]
    if file.filename == "" or file.filename is None:
        flash("Файл не обрано.", "error")
        return redirect(url_for("upload_page"))

    if not allowed_file(file.filename):
        flash("Недопустимий формат файлу. Дозволені: .xlsx, .xls, .csv, .docx", "error")
        return redirect(url_for("upload_page"))

    filename = secure_filename(file.filename)
    tmp_dir = tempfile.mkdtemp()
    filepath = os.path.join(tmp_dir, filename)
    file.save(filepath)

    result = parse_uploaded_file(filepath)

    try:
        os.remove(filepath)
        os.rmdir(tmp_dir)
    except OSError:
        pass

    if result.get("error"):
        flash(result["error"], "error")
        return redirect(url_for("upload_page"))

    session["parsed_data"] = json.dumps(result, ensure_ascii=False)

    return render_template(
        "preview.html",
        group_name=result["group_name"],
        students=result["students"],
        disciplines=result["disciplines"],
    )


@app.route("/import-gdoc", methods=["POST"])
def handle_gdoc_import():
    """Import and parse a Google Doc URL as a .docx file."""
    doc_url = request.form.get("doc_url", "").strip()
    if not doc_url:
        flash("Введіть посилання на Google Doc.", "error")
        return redirect(url_for("upload_page"))

    doc_id = extract_google_doc_id(doc_url)
    if not doc_id:
        flash("Неправильний формат посилання на Google Doc. Будь ласка, введіть коректне посилання.", "error")
        return redirect(url_for("upload_page"))

    tmp_dir = tempfile.mkdtemp()
    filename = f"{doc_id}.docx"
    filepath = os.path.join(tmp_dir, filename)

    try:
        import requests

        creds = get_drive_credentials()

        headers = {"Authorization": f"Bearer {creds.token}"}
        export_url = f"https://www.googleapis.com/drive/v3/files/{doc_id}/export?mimeType=application/vnd.openxmlformats-officedocument.wordprocessingml.document"

        response = requests.get(export_url, headers=headers)
        if response.status_code != 200:
            logger.error("Failed to export Google Doc: %s", response.text)
            flash(
                "Не вдалося завантажити Google Doc. Переконайтеся, що документ доступний для перегляду "
                "тому Google-акаунту, яким авторизований сервіс (сервісний акаунт або той, хто пройшов "
                "OAuth-вхід), або зробіть документ доступним за посиланням.",
                "error"
            )
            return redirect(url_for("upload_page"))

        with open(filepath, "wb") as f:
            f.write(response.content)

        result = parse_docx_file(filepath)

    except Exception as exc:
        logger.exception("Google Doc import failed")
        flash(f"Помилка при імпорті з Google Doc: {exc}", "error")
        return redirect(url_for("upload_page"))
    finally:
        try:
            if os.path.exists(filepath):
                os.remove(filepath)
            if os.path.exists(tmp_dir):
                os.rmdir(tmp_dir)
        except OSError:
            pass

    if result.get("error"):
        flash(result["error"], "error")
        return redirect(url_for("upload_page"))

    session["parsed_data"] = json.dumps(result, ensure_ascii=False)

    return render_template(
        "preview.html",
        group_name=result["group_name"],
        students=result["students"],
        disciplines=result["disciplines"],
    )


@app.route("/create-journal", methods=["POST"])
def create_journal():
    """Create a Google Sheets journal from previously parsed data."""
    raw = session.get("parsed_data")
    if not raw:
        flash("Дані не знайдено. Будь ласка, завантажте файл або Google Doc спочатку.", "error")
        return redirect(url_for("upload_page"))

    data = json.loads(raw)
    share_email = request.form.get("email", "").strip() or None
    folder_id = request.form.get("folder_id", "").strip() or None

    for idx, d in enumerate(data["disciplines"]):
        t_email = request.form.get(f"teacher_email_{idx}", "").strip() or None
        d["teacher_email"] = t_email

    try:
        result = create_academic_journal(
            group_name=data["group_name"],
            students=data["students"],
            disciplines=data["disciplines"],
            share_email=share_email,
            folder_id=folder_id,
        )
    except FileNotFoundError as exc:
        flash(str(exc), "error")
        return redirect(url_for("upload_page"))
    except Exception as exc:
        logger.exception("Failed to create Google Sheets journal")
        flash(f"Помилка при створенні Google Sheets: {exc}", "error")
        return redirect(url_for("upload_page"))

    session.pop("parsed_data", None)

    return render_template(
        "success.html",
        title=result["title"],
        spreadsheet_url=result["spreadsheet_url"],
        tab_count=len(data["disciplines"]) or 1,
        shared_with=share_email,
    )


@app.route("/upload-roster", methods=["POST"])
def upload_roster():
    """Receive the multi-group student roster (.doc/.docx) and persist it."""
    file = request.files.get("roster_file")
    if not file or file.filename == "":
        flash("Файл списку студентів не обрано.", "error")
        return redirect(url_for("upload_page"))

    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in ALLOWED_ROSTER_RNP_EXTENSIONS:
        flash("Список студентів має бути у форматі .doc або .docx.", "error")
        return redirect(url_for("upload_page"))

    filepath = _save_upload_to_tmp(file)
    try:
        result = parse_student_roster(filepath)
    finally:
        _cleanup_tmp(filepath)

    if result["error"]:
        flash(result["error"], "error")
        return redirect(url_for("upload_page"))

    save_roster(result["groups"])
    total_students = sum(len(v) for v in result["groups"].values())
    flash(
        f"Список студентів збережено: {len(result['groups'])} груп, {total_students} студентів.",
        "success",
    )
    return redirect(url_for("upload_page"))


@app.route("/upload-rnp", methods=["POST"])
def upload_rnp():
    """Receive an RNP (.doc/.docx) for one specialty, match it against the stored
    roster, and show a batch preview of every group it covers."""
    roster = load_roster()
    if not roster:
        flash("Спершу завантажте список студентів.", "error")
        return redirect(url_for("upload_page"))

    file = request.files.get("rnp_file")
    if not file or file.filename == "":
        flash("Файл РНП не обрано.", "error")
        return redirect(url_for("upload_page"))

    ext = file.filename.rsplit(".", 1)[-1].lower() if "." in file.filename else ""
    if ext not in ALLOWED_ROSTER_RNP_EXTENSIONS:
        flash("РНП має бути у форматі .doc або .docx.", "error")
        return redirect(url_for("upload_page"))

    filepath = _save_upload_to_tmp(file)
    try:
        rnp = parse_rnp(filepath)
    finally:
        _cleanup_tmp(filepath)

    if rnp["error"]:
        flash(rnp["error"], "error")
        return redirect(url_for("upload_page"))

    match = match_groups_to_rnp(roster, rnp)
    if not match["matched"]:
        flash(
            f"Жодна група зі збереженого списку не належить до спеціальності "
            f"«{rnp['specialty_code']}» з цього РНП.",
            "error",
        )
        return redirect(url_for("upload_page"))

    batch_id = uuid.uuid4().hex
    save_batch_state(batch_id, match)
    session["batch_id"] = batch_id

    return render_template(
        "preview_batch.html",
        specialty=rnp["specialty_code"],
        matched=match["matched"],
        skipped=match["skipped"],
    )


@app.route("/create-journals-batch", methods=["POST"])
def create_journals_batch():
    """Create one Google Sheets journal per matched group from the batch preview."""
    batch_id = session.get("batch_id")
    match = load_batch_state(batch_id) if batch_id else None
    if not match:
        flash("Дані пакету не знайдено або застаріли. Завантажте РНП ще раз.", "error")
        return redirect(url_for("upload_page"))

    share_email = request.form.get("email", "").strip() or None
    folder_id = request.form.get("folder_id", "").strip() or None

    results = []
    errors = []
    for group_name, info in match["matched"].items():
        disciplines = info["disciplines"]
        for idx, d in enumerate(disciplines):
            field = f"teacher_email_{group_name}_{idx}"
            d["teacher_email"] = request.form.get(field, "").strip() or None

        try:
            result = create_academic_journal(
                group_name=group_name,
                students=info["students"],
                disciplines=disciplines,
                share_email=share_email,
                folder_id=folder_id,
            )
            results.append({"group": group_name, **result})
        except Exception as exc:
            logger.exception("Failed to create journal for group %s", group_name)
            errors.append({"group": group_name, "error": str(exc)})

    delete_batch_state(batch_id)
    session.pop("batch_id", None)

    return render_template(
        "success_batch.html",
        results=results,
        errors=errors,
        shared_with=share_email,
    )


if __name__ == "__main__":
    app.run(debug=True, port=5000)
