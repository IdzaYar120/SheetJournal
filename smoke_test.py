"""End-to-end smoke test: upload sample_data.csv, generate a journal and
download the resulting .xlsx through the Flask routes. Used by CI before
building the desktop .exe.  Run: python smoke_test.py"""

import io
import os
import re
import tempfile

os.environ.setdefault("SHEETJOURNAL_DATA_DIR", tempfile.mkdtemp())

from openpyxl import load_workbook

from app import app


def main() -> None:
    client = app.test_client()
    with open("sample_data.csv", "rb") as f:
        resp = client.post("/upload", data={"file": (f, "sample_data.csv")},
                           content_type="multipart/form-data")
    assert resp.status_code == 200, resp.status_code

    resp = client.post("/create-journal", data={})
    assert resp.status_code == 200, resp.status_code
    match = re.search(r'/download/[0-9a-f]+', resp.get_data(as_text=True))
    assert match, "no download link on the success page"

    resp = client.get(match.group(0))
    assert resp.status_code == 200, resp.status_code
    wb = load_workbook(io.BytesIO(resp.data))
    print("OK:", len(wb.sheetnames), "sheets:", ", ".join(wb.sheetnames))


if __name__ == "__main__":
    main()
