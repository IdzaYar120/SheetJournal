"""
Academic Journal generation aligned with the reference workbook structure.

Builds a fully-formatted .xlsx file directly with openpyxl — no Google API
calls, no credentials, no consent screens. The workbook can then simply be
dragged into Google Drive (which converts it to a Google Sheet) or used as a
plain Excel file.

A small amount of Google-auth machinery is kept below only for the optional
"import from Google Docs" feature in app.py (reading a Google Doc the user
already has), which is unrelated to generating the journal itself.
"""

import json
import logging
import math
import datetime
import re
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)


# ==========================================================================
# Google auth (only used by app.py's optional Google Doc import feature)
# ==========================================================================

CREDENTIALS_PATH = Path(__file__).parent / "credentials.json"
AUTHORIZED_USER_PATH = Path(__file__).parent / "authorized_user.json"

OAUTH_SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive",
]


def _credentials_kind() -> str:
    """Classify CREDENTIALS_PATH as 'service_account', 'oauth_client', 'missing', or 'unknown'."""
    if not CREDENTIALS_PATH.exists():
        return "missing"
    try:
        with open(CREDENTIALS_PATH, "r", encoding="utf-8") as f:
            data = json.load(f)
    except Exception:
        return "unknown"
    if data.get("type") == "service_account":
        return "service_account"
    if "installed" in data or "web" in data:
        return "oauth_client"
    return "unknown"


def _try_application_default_credentials():
    """Try `gcloud`-issued Application Default Credentials, if present."""
    try:
        import google.auth
        creds, _project = google.auth.default(scopes=OAUTH_SCOPES)
        return creds
    except Exception:
        return None


def _get_client():
    """Authenticate and return a gspread client (used only for Google Doc import).

    Tries, in order: Application Default Credentials, a service-account key,
    or an OAuth Client ID (Desktop app) at CREDENTIALS_PATH.
    """
    import gspread

    adc = _try_application_default_credentials()
    if adc is not None:
        return gspread.authorize(adc)

    kind = _credentials_kind()
    if kind == "missing":
        raise FileNotFoundError(
            "Немає жодного способу авторизації в Google. Або виконайте "
            "'gcloud auth application-default login --scopes=...' (див. INSTRUCTION.md), "
            f"або покладіть credentials.json (ключ сервісного акаунту чи OAuth Client ID) за шляхом: {CREDENTIALS_PATH}"
        )
    if kind == "service_account":
        return gspread.service_account(filename=str(CREDENTIALS_PATH))
    if kind == "oauth_client":
        return gspread.oauth(
            scopes=OAUTH_SCOPES,
            credentials_filename=str(CREDENTIALS_PATH),
            authorized_user_filename=str(AUTHORIZED_USER_PATH),
        )
    raise RuntimeError(
        "Файл credentials.json має нерозпізнаний формат — це має бути або ключ "
        "сервісного акаунту (поле \"type\": \"service_account\"), або OAuth Client ID "
        "(Desktop app, поле \"installed\") у форматі JSON."
    )


def get_drive_credentials():
    """Return valid Google credentials for direct REST calls, e.g. exporting
    a Google Doc via the Drive API from app.py."""
    import google.auth.transport.requests

    creds = _get_client().auth
    if not creds.valid:
        creds.refresh(google.auth.transport.requests.Request())
    return creds


def get_service_account_email() -> str | None:
    """Return the service-account email for display, or None when using OAuth/ADC
    (where the signed-in Google account itself is the actor, not a robot email)."""
    if _credentials_kind() != "service_account":
        return None
    try:
        with open(CREDENTIALS_PATH, "r", encoding="utf-8") as f:
            return json.load(f).get("client_email")
    except Exception:
        return None


# ==========================================================================
# openpyxl helpers — thin wrappers so the sheet-building logic below reads
# the same as it would against a live spreadsheet API.
# ==========================================================================

_BORDER_SIDE = Side(style="thin", color="B3B3B3")
_THIN_BORDER = Border(left=_BORDER_SIDE, right=_BORDER_SIDE, top=_BORDER_SIDE, bottom=_BORDER_SIDE)


def _col_letter(col_index: int) -> str:
    """Convert a 1-based column index to a spreadsheet column letter (A, B, …, Z, AA, …)."""
    return get_column_letter(col_index)


def _hex(r: float, g: float, b: float) -> str:
    """Convert 0-1 float RGB (as used by the old Google Sheets format blocks) to hex."""
    return f"{round(r * 255):02X}{round(g * 255):02X}{round(b * 255):02X}"


def _coerce(val):
    """Numeric-looking strings become real numbers (so SUM/AVERAGE formulas work),
    formula strings ('=...') are left as-is for openpyxl to store as formulas."""
    if isinstance(val, str) and not val.startswith("="):
        try:
            return float(val)
        except ValueError:
            return val
    return val


def _write_rows(ws: Worksheet, rows_data: list[list], start_row: int = 1, start_col: int = 1) -> None:
    """Write a list-of-lists into a worksheet starting at start_row/start_col."""
    for r_off, row in enumerate(rows_data):
        for c_off, val in enumerate(row):
            if val is None or val == "":
                continue
            ws.cell(row=start_row + r_off, column=start_col + c_off).value = _coerce(val)


def _apply_format(ws: Worksheet, cell_range: str, *, bold=None, size=None, color=None, bg=None, halign=None, valign=None) -> None:
    font_kwargs = {}
    if bold is not None:
        font_kwargs["bold"] = bold
    if size is not None:
        font_kwargs["size"] = size
    if color is not None:
        font_kwargs["color"] = color
    font = Font(**font_kwargs) if font_kwargs else None
    fill = PatternFill(fill_type="solid", fgColor=bg) if bg else None
    align = Alignment(horizontal=halign, vertical=valign) if (halign or valign) else None
    for row in ws[cell_range]:
        for cell in row:
            if font is not None:
                cell.font = font
            if fill is not None:
                cell.fill = fill
            if align is not None:
                cell.alignment = align


def _apply_borders(ws: Worksheet, cell_range: str) -> None:
    for row in ws[cell_range]:
        for cell in row:
            cell.border = _THIN_BORDER


def _set_col_width(ws: Worksheet, start_col_letter: str, pixels: float, end_col_letter: str | None = None) -> None:
    width = max(2.0, round((pixels - 5) / 7, 1))
    start_idx = column_index_from_string(start_col_letter)
    end_idx = column_index_from_string(end_col_letter) if end_col_letter else start_idx
    for idx in range(start_idx, end_idx + 1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _freeze(ws: Worksheet, rows: int, cols: int) -> None:
    ws.freeze_panes = ws.cell(row=rows + 1, column=cols + 1).coordinate


# ==========================================================================
# Discipline naming / lookup helpers
# ==========================================================================

def get_short_name(name: str) -> str:
    """Generate a short Cyrillic abbreviation or truncated string for a discipline name."""
    clean = name.replace("Курсовий проект з", "").replace("Курсова робота з", "")
    clean = clean.replace("Курсовий проект", "").replace("Курсова робота", "")
    clean = clean.replace(",", " ").replace("-", " ").strip()

    stop_words = {"та", "і", "в", "на", "з", "за", "для", "по", "у", "й", "о"}
    words = [w for w in clean.split() if w.lower() not in stop_words]

    if not words:
        return "Дисц"

    if len(words) == 1:
        w = words[0]
        return w[:3].capitalize()

    acronym = ""
    for w in words:
        if w:
            acronym += w[0].upper()
    return acronym[:5]


def get_first_v_tab(disciplines: list[dict]) -> str:
    """Name of the first В_* report tab, the one other report tabs link names from."""
    for d in disciplines:
        if d.get("control_type") != "course_project":
            return f"В_{get_short_name(d['name'])}"
    return f"В_{get_short_name(disciplines[0]['name'])}_КР"


def get_sk_cols_for_discipline(d: dict, disciplines: list[dict]) -> dict:
    """Determine the column letters for a discipline in the СК worksheet."""
    col = 3
    for x in disciplines:
        if x is d:
            if x.get("control_type") == "course_project":
                return {
                    "def_points": _col_letter(col),
                    "def_pct": _col_letter(col + 1),
                    "cnt_points": _col_letter(col + 2),
                    "cnt_pct": _col_letter(col + 3)
                }
            else:
                return {
                    "points": _col_letter(col),
                    "pct": _col_letter(col + 1)
                }
        if x.get("control_type") == "course_project":
            col += 4
        else:
            col += 2
    return {}


# ==========================================================================
# Sheet builders
# ==========================================================================

def _populate_pc_sheet(
    worksheet: Worksheet,
    students: list[str],
    discipline_name: str,
    class_count: int,
    group_name: str,
    control_type: str = "credit",
    teacher_email: str | None = None,
    first_pc_tab_name: str | None = None,
) -> None:
    """Fill and format a single ПК_{ShortName} worksheet."""
    class_count = max(class_count, 1)
    total_cols = 2 + class_count
    total_rows = 4 + len(students) + 1

    rows_data: list[list] = []

    # Row 1: Title Card
    label = f"{group_name} — {discipline_name}"
    if control_type == "exam":
        label += " (Екзамен)"
    if teacher_email:
        label += f" (Викладач: {teacher_email})"
    row1 = [label] + [""] * (total_cols - 1)
    rows_data.append(row1)

    # Row 2: Date row
    start_date = datetime.date(2026, 2, 2)
    dates = []
    for col_idx in range(class_count):
        dates.append(start_date + datetime.timedelta(days=7 * col_idx))
    row2 = ["", ""] + dates
    rows_data.append(row2)

    # Row 3: Class type
    if "фізичн" in discipline_name.lower() or "іноземн" in discipline_name.lower():
        types = ["п"] * class_count
    else:
        types = []
        for i in range(class_count):
            types.append("л" if i % 2 == 0 else "п")
    row3 = ["вид занять", ""] + types
    rows_data.append(row3)

    # Row 4: Max points
    target_total = 80.0 if control_type == "credit" else 50.0
    num_practicals = sum(1 for t in types if t == "п")
    if num_practicals > 0:
        val = round(target_total / num_practicals, 1)
    else:
        val = 10.0
    max_points = [val if t == "п" else "" for t in types]
    row4 = ["бали", ""] + max_points
    rows_data.append(row4)

    # Student rows (Row 5+)
    for idx, name in enumerate(students):
        row_num = 5 + idx
        no_cell = f"{idx + 1}.0"

        if first_pc_tab_name:
            name_cell = f"='{first_pc_tab_name}'!B{row_num}"
        else:
            name_cell = name

        row = [no_cell, name_cell] + [""] * class_count
        rows_data.append(row)

    # Serial numbers row at the bottom
    serial_row = ["", ""] + [f"{i + 1}.0" for i in range(class_count)]
    rows_data.append(serial_row)

    _write_rows(worksheet, rows_data)

    last_col_letter = _col_letter(total_cols)

    _apply_format(worksheet, f"A1:{last_col_letter}1", bold=True, size=12, color=_hex(1.0, 1.0, 1.0),
                  bg=_hex(0.15, 0.27, 0.49), halign="center", valign="center")
    worksheet.merge_cells(f"A1:{last_col_letter}1")

    _apply_format(worksheet, f"A2:{last_col_letter}4", bold=True, size=9,
                  bg=_hex(0.93, 0.94, 0.96), halign="center", valign="center")

    _apply_format(worksheet, f"B5:B{4 + len(students)}", halign="left", valign="center")
    _apply_format(worksheet, f"A5:A{4 + len(students)}", halign="center", valign="center")

    if class_count > 0:
        class_start = _col_letter(3)
        _apply_format(worksheet, f"{class_start}5:{last_col_letter}{5 + len(students)}", halign="center", valign="center")

    _set_col_width(worksheet, "A", 45)
    _set_col_width(worksheet, "B", 260)
    if class_count > 0:
        _set_col_width(worksheet, class_start, 35, end_col_letter=last_col_letter)

    _apply_borders(worksheet, f"A1:{last_col_letter}{total_rows}")
    _freeze(worksheet, rows=4, cols=2)


def _populate_kr_sheets(
    ws_content: Worksheet,
    ws_defense: Worksheet,
    students: list[str],
    discipline_name: str,
    group_name: str,
    teacher_email: str | None = None,
    first_pc_tab_name: str | None = None,
) -> None:
    """Populate content and defense worksheets for course projects."""
    num_students = len(students)
    total_rows = 2 + 3 * num_students + 2

    c_title = ws_content.title

    row1_c = [
        f"{group_name} — Курсовий проект з дисципліни: {discipline_name} (Зміст)" + (f" (Викладач: {teacher_email})" if teacher_email else ""),
        "", "", "", "", "", "", "", ""
    ]
    row2_c = [
        "", "Підготовка курсової роботи (Максимальні бали)",
        "5.0", "5.0", "10.0", "10.0", "15.0", "15.0", "=SUM(C2:H2)"
    ]
    rows_c = [row1_c, row2_c]

    for idx, name in enumerate(students):
        row_num = 3 + idx
        name_cell = f"='{first_pc_tab_name}'!B{row_num + 2}" if first_pc_tab_name else name
        rows_c.append([
            f"{idx+1}.0", name_cell, "", "", "", "", "", "", f"=SUM(C{row_num}:H{row_num})"
        ])

    rows_c.append(["", "Комісія Член 2 (Максимальні бали)", "5.0", "5.0", "10.0", "10.0", "15.0", "15.0", f"=SUM(C{3+num_students}:H{3+num_students})"])

    for idx in range(num_students):
        row_num = 3 + num_students + 1 + idx
        rows_c.append([
            f"{idx+1}.0", f"=B{3+idx}", "", "", "", "", "", "", f"=SUM(C{row_num}:H{row_num})"
        ])

    rows_c.append(["", "Комісія Член 3 (Максимальні бали)", "5.0", "5.0", "10.0", "10.0", "15.0", "15.0", f"=SUM(C{3+2*num_students+1}:H{3+2*num_students+1})"])

    for idx in range(num_students):
        row_num = 3 + 2 * num_students + 2 + idx
        rows_c.append([
            f"{idx+1}.0", f"=B{3+num_students+1+idx}", "", "", "", "", "", "", f"=SUM(C{row_num}:H{row_num})"
        ])

    _write_rows(ws_content, rows_c)

    row1_d = [
        f"{group_name} — Курсовий проект з дисципліни: {discipline_name} (Захист)" + (f" (Викладач: {teacher_email})" if teacher_email else ""),
        "", "", "", "", ""
    ]
    row2_d = [
        "", "Захист курсової роботи (Максимальні бали)",
        "15.0", "10.0", "15.0", "=SUM(C2:E2)"
    ]
    rows_d = [row1_d, row2_d]

    for idx, name in enumerate(students):
        row_num = 3 + idx
        rows_d.append([
            f"{idx+1}.0", f"='{c_title}'!B{row_num}", "", "", "", f"=SUM(C{row_num}:E{row_num})"
        ])

    rows_d.append(["", "Комісія Член 2 (Максимальні бали)", "15.0", "10.0", "15.0", f"=SUM(C{3+num_students}:E{3+num_students})"])

    for idx in range(num_students):
        row_num = 3 + num_students + 1 + idx
        rows_d.append([
            f"{idx+1}.0", f"=B{3+idx}", "", "", "", f"=SUM(C{row_num}:E{row_num})"
        ])

    rows_d.append(["", "Комісія Член 3 (Максимальні бали)", "15.0", "10.0", "15.0", f"=SUM(C{3+2*num_students+1}:E{3+2*num_students+1})"])

    for idx in range(num_students):
        row_num = 3 + 2 * num_students + 2 + idx
        rows_d.append([
            f"{idx+1}.0", f"=B{3+num_students+1+idx}", "", "", "", f"=SUM(C{row_num}:E{row_num})"
        ])

    _write_rows(ws_defense, rows_d)

    spacer1_row = 3 + num_students
    spacer2_row = 4 + 2 * num_students

    for ws, cols_count in [(ws_content, 9), (ws_defense, 6)]:
        last_col = _col_letter(cols_count)
        _apply_format(ws, f"A1:{last_col}1", bold=True, size=12, color=_hex(1.0, 1.0, 1.0),
                      bg=_hex(0.45, 0.15, 0.20), halign="center", valign="center")
        _apply_format(ws, f"A2:{last_col}2", bold=True, size=9,
                      bg=_hex(0.95, 0.9, 0.9), halign="center", valign="center")
        _apply_format(ws, f"A{spacer1_row}:{last_col}{spacer1_row}", bold=True, size=9,
                      bg=_hex(0.95, 0.9, 0.9), halign="center", valign="center")
        _apply_format(ws, f"A{spacer2_row}:{last_col}{spacer2_row}", bold=True, size=9,
                      bg=_hex(0.95, 0.9, 0.9), halign="center", valign="center")

        ws.merge_cells(f"A1:{last_col}1")
        _set_col_width(ws, "A", 45)
        _set_col_width(ws, "B", 260)
        _apply_borders(ws, f"A1:{last_col}{total_rows}")
        _freeze(ws, rows=2, cols=2)


def _create_milestone_sheet(
    worksheet: Worksheet,
    students: list[str],
    disciplines: list[dict],
    milestone: int,
    group_name: str,
    first_pc_tab_name: str | None = None,
) -> Worksheet:
    """Fill and format the РК1/РК2 worksheet."""
    milestone_disciplines = [d for d in disciplines if d.get("control_type", "credit") != "course_project"]
    if not milestone_disciplines:
        milestone_disciplines = [{"name": group_name, "class_count": 0, "control_type": "credit"}]

    num_students = len(students)
    total_cols = 2 + len(milestone_disciplines) * 3
    total_rows = 2 + num_students

    rows_data: list[list] = []

    row1 = ["№ з/п", "ПІП"]
    for d in milestone_disciplines:
        short = get_short_name(d["name"])
        pc_tab = f"ПК_{short}"
        row1.extend([f"='{pc_tab}'!C1", "", ""])
    rows_data.append(row1)

    row2 = ["", ""]
    for d in milestone_disciplines:
        short = get_short_name(d["name"])
        pc_tab = f"ПК_{short}"
        class_count = max(d.get("class_count") or 0, 1)
        middle = math.ceil(class_count / 2)

        if milestone == 1:
            start_col = "C"
            end_col = _col_letter(2 + middle)
        else:
            start_col = _col_letter(3 + middle)
            end_col = _col_letter(2 + class_count)

        sum_formula = f"=SUM('{pc_tab}'!{start_col}4:{end_col}4)"
        row2.extend([sum_formula, "%", "Національна / ECTS"])
    rows_data.append(row2)

    for idx in range(num_students):
        row_num = 3 + idx
        student_no = f"{idx+1}.0"
        name_cell = f"='{first_pc_tab_name}'!B{row_num + 2}" if first_pc_tab_name else students[idx]
        row = [student_no, name_cell]

        col_idx = 3
        for d in milestone_disciplines:
            short = get_short_name(d["name"])
            pc_tab = f"ПК_{short}"
            class_count = max(d.get("class_count") or 0, 1)
            middle = math.ceil(class_count / 2)

            if milestone == 1:
                start_col = "C"
                end_col = _col_letter(2 + middle)
            else:
                start_col = _col_letter(3 + middle)
                end_col = _col_letter(2 + class_count)

            pts_letter = _col_letter(col_idx)
            pct_letter = _col_letter(col_idx + 1)

            pts_formula = f"=SUM('{pc_tab}'!{start_col}{row_num+2}:{end_col}{row_num+2})"
            pct_formula = f"={pts_letter}{row_num}/{pts_letter}$2*100"
            ects_formula = (
                f'=IF(AND({pct_letter}{row_num}<=100,88<={pct_letter}{row_num}),"5A",'
                f'IF(AND({pct_letter}{row_num}<88,81<={pct_letter}{row_num}),"4B",'
                f'IF(AND({pct_letter}{row_num}<81,74<={pct_letter}{row_num}),"4C",'
                f'IF(AND({pct_letter}{row_num}<74,67<={pct_letter}{row_num}),"3D",'
                f'IF(AND({pct_letter}{row_num}<67,60<={pct_letter}{row_num}),"3E",'
                f'IF(AND({pct_letter}{row_num}<60,35<={pct_letter}{row_num}),"2FX",'
                f'IF(AND({pct_letter}{row_num}<35,0<={pct_letter}{row_num}),"2F","")))))))'
            )
            row.extend([pts_formula, pct_formula, ects_formula])
            col_idx += 3

        rows_data.append(row)

    _write_rows(worksheet, rows_data)

    last_col_letter = _col_letter(total_cols)
    _apply_format(worksheet, f"A1:{last_col_letter}2", bold=True, size=9,
                  bg=_hex(0.88, 0.92, 0.95), halign="center", valign="center")
    _set_col_width(worksheet, "A", 45)
    _set_col_width(worksheet, "B", 260)

    for i in range(len(milestone_disciplines)):
        sub_col = 3 + i * 3
        start_letter = _col_letter(sub_col)
        end_letter = _col_letter(sub_col + 2)
        worksheet.merge_cells(f"{start_letter}1:{end_letter}1")

    _apply_borders(worksheet, f"A1:{last_col_letter}{total_rows}")
    _freeze(worksheet, rows=2, cols=2)
    return worksheet


def _create_helper_sheet(
    worksheet: Worksheet,
    students: list[str],
    disciplines: list[dict],
    title: str,
    ref_sheet: str | None = None,
    first_pc_tab_name: str | None = None,
    group_name: str = "",
) -> Worksheet:
    """Fill and format helper sheets like ДОД_РК1, ДОД_РК2, ІНДЗ, ДОД_ІНДЗ."""
    milestone_disciplines = [d for d in disciplines if d.get("control_type", "credit") != "course_project"]
    if not milestone_disciplines:
        milestone_disciplines = [{"name": group_name, "class_count": 0, "control_type": "credit"}]

    num_students = len(students)
    total_cols = 2 + len(milestone_disciplines) * 2
    total_rows = 2 + num_students

    rows_data: list[list] = []

    row1 = ["№ з/п", "ПІП"]
    for d in milestone_disciplines:
        short = get_short_name(d["name"])
        pc_tab = f"ПК_{short}"
        row1.extend([f"='{pc_tab}'!C1", ""])
    rows_data.append(row1)

    row2 = ["", ""]
    for d in milestone_disciplines:
        max_val = "10.0" if title == "ІНДЗ" else ""
        row2.extend([max_val, "%"])
    rows_data.append(row2)

    for idx in range(num_students):
        row_num = 3 + idx
        student_no = f"{idx+1}.0"
        name_cell = f"='{first_pc_tab_name}'!B{row_num + 2}" if first_pc_tab_name else students[idx]
        row = [student_no, name_cell]

        for i, d in enumerate(milestone_disciplines):
            col_idx = 3 + i * 2
            pts_letter = _col_letter(col_idx)

            if title == "ІНДЗ":
                pct_formula = f"={pts_letter}{row_num}/${pts_letter}$2*100"
            elif title == "ДОД_ІНДЗ":
                pct_formula = f"={pts_letter}{row_num}/'ІНДЗ'!${pts_letter}$2*100"
            else:
                ref_pts_col = _col_letter(3 + i * 3)
                pct_formula = f"={pts_letter}{row_num}/'{ref_sheet}'!${ref_pts_col}$2*100"

            row.extend(["", pct_formula])
        rows_data.append(row)

    _write_rows(worksheet, rows_data)

    last_col_letter = _col_letter(total_cols)
    _apply_format(worksheet, f"A1:{last_col_letter}2", bold=True, size=9,
                  bg=_hex(0.9, 0.94, 0.92), halign="center", valign="center")
    _set_col_width(worksheet, "A", 45)
    _set_col_width(worksheet, "B", 260)

    for i in range(len(milestone_disciplines)):
        sub_col = 3 + i * 2
        start_letter = _col_letter(sub_col)
        end_letter = _col_letter(sub_col + 1)
        worksheet.merge_cells(f"{start_letter}1:{end_letter}1")

    _apply_borders(worksheet, f"A1:{last_col_letter}{total_rows}")
    _freeze(worksheet, rows=2, cols=2)
    return worksheet


def _create_sk_sheet(
    worksheet: Worksheet,
    students: list[str],
    disciplines: list[dict],
    group_name: str,
    first_pc_tab_name: str | None = None,
) -> Worksheet:
    """Fill and format the Semester Control worksheet ("СК")."""
    num_students = len(students)

    total_cols = 2
    for d in disciplines:
        if d.get("control_type", "credit") == "course_project":
            total_cols += 4
        else:
            total_cols += 2

    rows_data: list[list] = []

    row1 = ["№ з/п", "ПІП"]
    for d in disciplines:
        short = get_short_name(d["name"])
        if d.get("control_type", "credit") == "course_project":
            row1.extend([f"{d['name']} (курсова робота)", "", "", ""])
        else:
            pc_tab = f"ПК_{short}"
            row1.extend([f"='{pc_tab}'!C1", ""])
    rows_data.append(row1)

    row2 = ["", ""]
    milestone_disciplines = [d for d in disciplines if d.get("control_type", "credit") != "course_project"]

    for d in disciplines:
        if d.get("control_type", "credit") == "course_project":
            row2.extend(["40.0", "%", "60.0", "%"])
        else:
            m_idx = milestone_disciplines.index(d)
            rk_col = _col_letter(3 + m_idx * 3)
            dod_col = _col_letter(3 + m_idx * 2)
            indz_col = _col_letter(3 + m_idx * 2)

            max_formula = (
                f"='РК1'!{rk_col}2+'РК2'!{rk_col}2+"
                f"'ДОД_РК1'!{dod_col}2+'ДОД_РК2'!{dod_col}2+"
                f"'ІНДЗ'!{indz_col}2+'ДОД_ІНДЗ'!{indz_col}2"
            )
            row2.extend([max_formula, "%"])
    rows_data.append(row2)

    for idx in range(num_students):
        row_num = 3 + idx
        student_no = f"{idx+1}.0"
        name_cell = f"='{first_pc_tab_name}'!B{row_num + 2}" if first_pc_tab_name else students[idx]
        row = [student_no, name_cell]

        for d in disciplines:
            if d.get("control_type", "credit") == "course_project":
                short = get_short_name(d["name"])
                r1 = row_num
                r2 = row_num + num_students + 1
                r3 = row_num + 2 * num_students + 2

                def_formula = f"=AVERAGE('КР_{short}_Захист'!F{r1},'КР_{short}_Захист'!F{r2},'КР_{short}_Захист'!F{r3})"
                def_pct = f"={_col_letter(len(row)+1)}{row_num}/40.0*100"
                cnt_formula = f"=AVERAGE('КР_{short}_Зміст'!I{r1},'КР_{short}_Зміст'!I{r2},'КР_{short}_Зміст'!I{r3})"
                cnt_pct = f"={_col_letter(len(row)+3)}{row_num}/60.0*100"
                row.extend([def_formula, def_pct, cnt_formula, cnt_pct])
            else:
                m_idx = milestone_disciplines.index(d)
                rk_col = _col_letter(3 + m_idx * 3)
                dod_col = _col_letter(3 + m_idx * 2)
                indz_col = _col_letter(3 + m_idx * 2)

                pts_formula = (
                    f"='РК1'!{rk_col}{row_num}+'ДОД_РК1'!{dod_col}{row_num}+"
                    f"'РК2'!{rk_col}{row_num}+'ДОД_РК2'!{dod_col}{row_num}+"
                    f"'ІНДЗ'!{indz_col}{row_num}+'ДОД_ІНДЗ'!{indz_col}{row_num}"
                )
                pts_letter = _col_letter(len(row) + 1)
                pct_formula = f"={pts_letter}{row_num}/{pts_letter}$2*100"
                row.extend([pts_formula, pct_formula])

        rows_data.append(row)

    _write_rows(worksheet, rows_data)

    last_col_letter = _col_letter(total_cols)
    _apply_format(worksheet, f"A1:{last_col_letter}2", bold=True, size=9,
                  bg=_hex(0.85, 0.9, 0.88), halign="center", valign="center")
    _set_col_width(worksheet, "A", 45)
    _set_col_width(worksheet, "B", 260)

    curr_col = 3
    for d in disciplines:
        if d.get("control_type", "credit") == "course_project":
            start_letter = _col_letter(curr_col)
            end_letter = _col_letter(curr_col + 3)
            worksheet.merge_cells(f"{start_letter}1:{end_letter}1")
            curr_col += 4
        else:
            start_letter = _col_letter(curr_col)
            end_letter = _col_letter(curr_col + 1)
            worksheet.merge_cells(f"{start_letter}1:{end_letter}1")
            curr_col += 2

    _apply_borders(worksheet, f"A1:{last_col_letter}{2 + num_students}")
    _freeze(worksheet, rows=2, cols=2)
    return worksheet


def _create_v_report_sheet(
    worksheet: Worksheet,
    students: list[str],
    discipline: dict,
    disciplines: list[dict],
    group_name: str,
    first_pc_tab_name: str | None = None,
) -> Worksheet:
    """Fill and format a single official report sheet "В_{ShortName}" for a discipline."""
    short = get_short_name(discipline["name"])
    title = worksheet.title
    control_type = discipline.get("control_type", "credit")

    num_students = len(students)
    total_cols = 11

    rows_data: list[list] = []

    rows_data.append(["Університет економіки і підприємництва"] + [""] * 10)
    rows_data.append([""] * 11)

    first_v_tab = get_first_v_tab(disciplines)
    rows_data.append(["2025 / 2026 навчальний рік", "", "", "", "", "", "", "Група", f"='{first_v_tab}'!I3" if first_pc_tab_name and title != first_v_tab else group_name, "", ""])
    rows_data.append(["", "", "", "", "", "", "", "Курс", f"='{first_v_tab}'!I4" if first_pc_tab_name and title != first_v_tab else "I", "", ""])
    rows_data.append(["Заліково-екзаменаційна відомість № ", "", "", "", "", "", "", "", "", "", ""])

    pc_tab = f"ПК_{short}"
    rows_data.append(["з дисципліни", "", f"='{pc_tab}'!C1" if control_type != "course_project" else f"='КР_{short}_Зміст'!B1", "", "", "", "", "", "", "", ""])
    rows_data.append([""] * 11)
    rows_data.append(["за 2 навчальний семестр", "", "", "", "", "", "", "Екзамен / Залік", "", "", ""])
    rows_data.append(["Екзаменатор", "", "", "", "", "", "", "", "", "", ""])
    rows_data.append(["Викладачі, які ведуть практичні, семінарські і лабораторні роботи", "", "", "", "", "", "", "", "", "", ""])
    rows_data.append([""] * 11)

    if control_type == "course_project":
        header_row12 = ["№ з/п", "Прізвище, ініціали студента", "Номер індивід. навч. плану", "Оцінка курсової роботи", "", "", "", "", "", "Підпис керівника та членів комісії", ""]
        header_row13 = ["", "", "", "Національна", "", "ECTS", "", "Бали", "", "", ""]
    else:
        header_row12 = ["№ з/п", "Прізвище, ініціали студента", "Номер індивід. навч. плану", "Залікова оцінка", "", "", "Підпис   викла- дача", "Екзаменаційна оцінка", "", "", "Підпис екзаме- натора"]
        header_row13 = ["", "", "", "Національна", "ECTS", "Бали", "", "Національна", "ECTS", "Бали", ""]

    rows_data.append(header_row12)
    rows_data.append(header_row13)

    sk_cols = get_sk_cols_for_discipline(discipline, disciplines)

    for idx in range(num_students):
        row_num = 14 + idx
        sk_row = 3 + idx

        first_report_sheet = get_first_v_tab(disciplines)
        if title == first_report_sheet:
            name_formula = (
                f'=IFERROR(CONCATENATE(LEFT(\'СК\'!B{sk_row},(FIND(" ",\'СК\'!B{sk_row})-1))," ",'
                f'MID(\'СК\'!B{sk_row},FIND(" ",\'СК\'!B{sk_row})+1,1),". ",'
                f'MID(\'СК\'!B{sk_row},FIND(" ",\'СК\'!B{sk_row},FIND(" ",\'СК\'!B{sk_row})+1)+1,1),"."),\'СК\'!B{sk_row})'
            )
        else:
            name_formula = f"='{first_report_sheet}'!B{row_num}"

        no_cell = f"='{first_report_sheet}'!A{row_num}" if title != first_report_sheet else f"{idx+1}.0"
        record_book_no = f"='{first_report_sheet}'!C{row_num}" if title != first_report_sheet else f"{4000+idx+1}"

        if control_type == "course_project":
            pts_formula = f"='СК'!{sk_cols['def_points']}{sk_row}+'СК'!{sk_cols['cnt_points']}{sk_row}"
            nat_formula = (
                f'=IF(AND(H{row_num}<=100,88<=H{row_num}),"Відмінно",'
                f'IF(AND(H{row_num}<88,81<=H{row_num}),"Добре",'
                f'IF(AND(H{row_num}<81,74<=H{row_num}),"Добре",'
                f'IF(AND(H{row_num}<74,67<=H{row_num}),"Задовільно",'
                f'IF(AND(H{row_num}<67,60<=H{row_num}),"Задовільно",'
                f'IF(AND(H{row_num}<60,35<=H{row_num}),"Незадовільно",'
                f'IF(AND(H{row_num}<35,0<=H{row_num}),"Незадовільно","")))))))'
            )
            ects_formula = (
                f'=IF(AND(H{row_num}<=100,88<=H{row_num}),"A",'
                f'IF(AND(H{row_num}<88,81<=H{row_num}),"B",'
                f'IF(AND(H{row_num}<81,74<=H{row_num}),"C",'
                f'IF(AND(H{row_num}<74,67<=H{row_num}),"D",'
                f'IF(AND(H{row_num}<67,60<=H{row_num}),"E",'
                f'IF(AND(H{row_num}<60,35<=H{row_num}),"FX",'
                f'IF(AND(H{row_num}<35,0<=H{row_num}),"F","")))))))'
            )
            student_row = [no_cell, name_formula, record_book_no, nat_formula, "", ects_formula, "", pts_formula, "", "", ""]

        elif control_type == "exam":
            pts_formula = f"='СК'!{sk_cols['points']}{sk_row}"
            nat_formula = (
                f'=IF(AND(J{row_num}<=100,88<=J{row_num}),"Відмінно",'
                f'IF(AND(J{row_num}<88,81<=J{row_num}),"Добре",'
                f'IF(AND(J{row_num}<81,74<=J{row_num}),"Добре",'
                f'IF(AND(J{row_num}<74,67<=J{row_num}),"Задовільно",'
                f'IF(AND(J{row_num}<67,60<=J{row_num}),"Задовільно",'
                f'IF(AND(J{row_num}<60,35<=J{row_num}),"Незадовільно",'
                f'IF(AND(J{row_num}<35,0<=J{row_num}),"Незадовільно","")))))))'
            )
            ects_formula = (
                f'=IF(AND(J{row_num}<=100,88<=J{row_num}),"A",'
                f'IF(AND(J{row_num}<88,81<=J{row_num}),"B",'
                f'IF(AND(J{row_num}<81,74<=J{row_num}),"C",'
                f'IF(AND(J{row_num}<74,67<=J{row_num}),"D",'
                f'IF(AND(J{row_num}<67,60<=J{row_num}),"E",'
                f'IF(AND(J{row_num}<60,35<=J{row_num}),"FX",'
                f'IF(AND(J{row_num}<35,0<=J{row_num}),"F","")))))))'
            )
            student_row = [no_cell, name_formula, record_book_no, "", "", "", "", nat_formula, ects_formula, pts_formula, ""]

        else:  # credit
            pts_formula = f"='СК'!{sk_cols['points']}{sk_row}"
            nat_formula = (
                f'=IF(AND(F{row_num}<=100,60<=F{row_num}),"Зараховано",'
                f'IF(AND(F{row_num}<60,35<=F{row_num}),"Не зараховано",'
                f'IF(AND(F{row_num}<35,0<=F{row_num}),"Не зараховано","")))'
            )
            ects_formula = (
                f'=IF(AND(F{row_num}<=100,88<=F{row_num}),"A",'
                f'IF(AND(F{row_num}<88,81<=F{row_num}),"B",'
                f'IF(AND(F{row_num}<81,74<=F{row_num}),"C",'
                f'IF(AND(F{row_num}<74,67<=F{row_num}),"D",'
                f'IF(AND(F{row_num}<67,60<=F{row_num}),"E",'
                f'IF(AND(F{row_num}<60,35<=F{row_num}),"FX",'
                f'IF(AND(F{row_num}<35,0<=F{row_num}),"F","")))))))'
            )
            student_row = [no_cell, name_formula, record_book_no, nat_formula, ects_formula, pts_formula, "", "", "", "", ""]

        rows_data.append(student_row)

    last_student_row = 13 + num_students
    rows_data.append([""] * 11)

    r_grp_size = last_student_row + 2
    r_present = last_student_row + 3
    r_excellent = last_student_row + 4

    ects_col = "F" if control_type == "course_project" else ("I" if control_type == "exam" else "E")

    rows_data.append(["", "Студентів в групі", f"=A{last_student_row}", "", "", "", "Не з’явилось", "", "", f"=C{r_grp_size}-J{r_present}", ""])
    rows_data.append(["", "", "", "", "", "", "Було присутніх", "", "", f"=SUM(J{r_excellent}:J{r_excellent+6})", ""])
    rows_data.append(["", "Не допущено", "0.0", "", "", "", "Відмінно / А", "", "", f'=COUNTIF(${ects_col}$14:${ects_col}${last_student_row},"A")', ""])
    rows_data.append(["", "", "", "", "", "", "Добре / B", "", "", f'=COUNTIF(${ects_col}$14:${ects_col}${last_student_row},"B")', ""])
    rows_data.append(["", "", "", "", "", "", "Добре / C", "", "", f'=COUNTIF(${ects_col}$14:${ects_col}${last_student_row},"C")', ""])
    rows_data.append(["", "", "", "", "", "", "Задовільно / D", "", "", f'=COUNTIF(${ects_col}$14:${ects_col}${last_student_row},"D")', ""])
    rows_data.append(["", "", "", "", "", "", "Задовільно / E", "", "", f'=COUNTIF(${ects_col}$14:${ects_col}${last_student_row},"E")', ""])
    rows_data.append(["", "", "", "", "", "", "Незадовільно / FX", "", "", f'=COUNTIF(${ects_col}$14:${ects_col}${last_student_row},"FX")', ""])
    rows_data.append(["", "", "", "", "", "", "Незадовільно / F", "", "", f'=COUNTIF(${ects_col}$14:${ects_col}${last_student_row},"F")', ""])

    _write_rows(worksheet, rows_data)

    _apply_format(worksheet, "A12:K13", bold=True, size=9, bg=_hex(0.94, 0.94, 0.94), halign="center", valign="center")
    _set_col_width(worksheet, "A", 45)
    _set_col_width(worksheet, "B", 240)
    _set_col_width(worksheet, "C", 100)

    if control_type == "course_project":
        worksheet.merge_cells("D12:E12")
        worksheet.merge_cells("F12:G12")
        worksheet.merge_cells("H12:I12")
    else:
        worksheet.merge_cells("D12:F12")
        worksheet.merge_cells("H12:J12")

    _apply_borders(worksheet, f"A12:K{last_student_row}")
    _freeze(worksheet, rows=13, cols=2)
    return worksheet


def _create_consolidated_dashboard(
    worksheet: Worksheet,
    students: list[str],
    disciplines: list[dict],
    group_name: str,
    first_pc_tab_name: str | None = None,
) -> Worksheet:
    """Fill and format the final consolidated dashboard "Відомості"."""
    num_students = len(students)
    total_cols = 3 + len(disciplines) * 3
    total_rows = 10 + num_students

    rows_data: list[list] = []

    rows_data.append(["Університет економіки і підприємництва"] + [""] * (total_cols - 1))
    rows_data.append([""] * total_cols)
    rows_data.append(["", "", "", "", "", "", "", "ХІД"] + [""] * (total_cols - 8))
    rows_data.append(["", "складання заліків та екзаменів сесії", "", "", "", "", "", ""] + [""] * (total_cols - 8))
    rows_data.append(["", "2025 / 2026 навчальний рік", "", "", "", "", "", ""] + [""] * (total_cols - 8))

    first_report_sheet = get_first_v_tab(disciplines)
    rows_data.append(["", f"='{first_report_sheet}'!A8", "", "", "", "", "група", "", "", "", "", "", group_name] + [""] * (total_cols - 13))
    rows_data.append([""] * total_cols)

    row8 = ["№ з/п", "Прізвище, ініціали студента", "Номер індивід. навч. плану", "Заліки, Екзамени"] + [""] * (total_cols - 4)
    rows_data.append(row8)

    row9 = ["", "", ""]
    for d in disciplines:
        ctype = d.get("control_type", "credit")
        label = "КП" if ctype == "course_project" else ("Е" if ctype == "exam" else "З")
        row9.extend([label, "", ""])
    rows_data.append(row9)

    row10 = ["", "", ""]
    for d in disciplines:
        row10.extend([d["name"], "", ""])
    rows_data.append(row10)

    for idx in range(num_students):
        row_num = 11 + idx
        v_row = 14 + idx

        no_cell = f"='{first_report_sheet}'!A{v_row}"
        name_cell = f"='{first_report_sheet}'!B{v_row}"
        record_book = f"='{first_report_sheet}'!C{v_row}"
        row = [no_cell, name_cell, record_book]

        for i, d in enumerate(disciplines):
            short = get_short_name(d["name"])
            v_tab = f"В_{short}" if d.get("control_type") != "course_project" else f"В_{short}_КР"
            ctype = d.get("control_type", "credit")

            if ctype == "course_project":
                ects_col = "F"
                score_col = "H"
            elif ctype == "exam":
                ects_col = "I"
                score_col = "J"
            else:
                ects_col = "E"
                score_col = "F"

            ects_dash_col = _col_letter(3 + i * 3 + 2)

            nat_formula = (
                f'=IFS({ects_dash_col}{row_num}="A",5,{ects_dash_col}{row_num}="B",4,'
                f'{ects_dash_col}{row_num}="C",4,{ects_dash_col}{row_num}="D",3,'
                f'{ects_dash_col}{row_num}="E",3,{ects_dash_col}{row_num}="F"," ",'
                f'{ects_dash_col}{row_num}="FX"," ")'
            )
            ects_formula = f"='{v_tab}'!{ects_col}{v_row}"
            score_formula = f"='{v_tab}'!{score_col}{v_row}"
            row.extend([nat_formula, ects_formula, score_formula])

        rows_data.append(row)

    _write_rows(worksheet, rows_data)

    last_col_letter = _col_letter(total_cols)
    _apply_format(worksheet, f"A8:{last_col_letter}10", bold=True, size=9,
                  bg=_hex(0.85, 0.88, 0.93), halign="center", valign="center")
    _set_col_width(worksheet, "A", 45)
    _set_col_width(worksheet, "B", 240)
    _set_col_width(worksheet, "C", 100)

    worksheet.merge_cells(f"D8:{last_col_letter}8")

    for i in range(len(disciplines)):
        sub_col = 4 + i * 3
        start_letter = _col_letter(sub_col)
        end_letter = _col_letter(sub_col + 2)
        worksheet.merge_cells(f"{start_letter}9:{end_letter}9")
        worksheet.merge_cells(f"{start_letter}10:{end_letter}10")

    _apply_borders(worksheet, f"A8:{last_col_letter}{total_rows}")
    _freeze(worksheet, rows=10, cols=2)
    return worksheet


# ==========================================================================
# Orchestrator
# ==========================================================================

def create_academic_journal(
    group_name: str,
    students: list[str],
    disciplines: list[dict],
    output_dir: str | Path | None = None,
) -> dict:
    """Build a fully-formatted .xlsx academic journal workbook matching the
    reference structure, and save it to disk. Returns title/file_path/filename."""
    wb = Workbook()
    wb.remove(wb.active)

    sheets_map: dict[str, Worksheet] = {}
    first_pc_tab_name = None

    # A. ПК_* and КР_* sheets
    for disc in disciplines:
        disc_name = disc["name"]
        ctype = disc.get("control_type", "credit")
        short = get_short_name(disc_name)

        if ctype == "course_project":
            c_title = f"КР_{short}_Зміст"
            d_title = f"КР_{short}_Захист"
            sheets_map[c_title] = wb.create_sheet(title=c_title)
            sheets_map[d_title] = wb.create_sheet(title=d_title)
            if first_pc_tab_name is None:
                first_pc_tab_name = c_title
        else:
            pc_title = f"ПК_{short}"
            sheets_map[pc_title] = wb.create_sheet(title=pc_title)
            if first_pc_tab_name is None:
                first_pc_tab_name = pc_title

    # B. Milestone & helper sheets
    for name in ("РК1", "ДОД_РК1", "РК2", "ДОД_РК2", "ІНДЗ", "ДОД_ІНДЗ"):
        sheets_map[name] = wb.create_sheet(title=name)

    # C. СК sheet
    sheets_map["СК"] = wb.create_sheet(title="СК")

    # D. В_* report sheets
    for disc in disciplines:
        short = get_short_name(disc["name"])
        ctype = disc.get("control_type", "credit")
        v_title = f"В_{short}" if ctype != "course_project" else f"В_{short}_КР"
        sheets_map[v_title] = wb.create_sheet(title=v_title)

    # E. Consolidated Dashboard sheet
    sheets_map["Відомості"] = wb.create_sheet(title="Відомості")

    # -------------------------------------------------------------
    # Populate everything, in the same order as the sheets appear
    # -------------------------------------------------------------
    for disc in disciplines:
        disc_name = disc["name"]
        ctype = disc.get("control_type", "credit")
        short = get_short_name(disc_name)
        teacher_email = disc.get("teacher_email")

        if ctype == "course_project":
            c_title = f"КР_{short}_Зміст"
            d_title = f"КР_{short}_Захист"
            _populate_kr_sheets(
                ws_content=sheets_map[c_title],
                ws_defense=sheets_map[d_title],
                students=students,
                discipline_name=disc_name,
                group_name=group_name,
                teacher_email=teacher_email,
                first_pc_tab_name=first_pc_tab_name,
            )
        else:
            pc_title = f"ПК_{short}"
            class_count = disc.get("class_count") or 0
            _populate_pc_sheet(
                worksheet=sheets_map[pc_title],
                students=students,
                discipline_name=disc_name,
                class_count=class_count,
                group_name=group_name,
                control_type=ctype,
                teacher_email=teacher_email,
                first_pc_tab_name=None if pc_title == first_pc_tab_name else first_pc_tab_name,
            )

    _create_milestone_sheet(sheets_map["РК1"], students, disciplines, 1, group_name, first_pc_tab_name)
    _create_helper_sheet(sheets_map["ДОД_РК1"], students, disciplines, "ДОД_РК1", "РК1", first_pc_tab_name, group_name)
    _create_milestone_sheet(sheets_map["РК2"], students, disciplines, 2, group_name, first_pc_tab_name)
    _create_helper_sheet(sheets_map["ДОД_РК2"], students, disciplines, "ДОД_РК2", "РК2", first_pc_tab_name, group_name)
    _create_helper_sheet(sheets_map["ІНДЗ"], students, disciplines, "ІНДЗ", None, first_pc_tab_name, group_name)
    _create_helper_sheet(sheets_map["ДОД_ІНДЗ"], students, disciplines, "ДОД_ІНДЗ", "ІНДЗ", first_pc_tab_name, group_name)

    _create_sk_sheet(sheets_map["СК"], students, disciplines, group_name, first_pc_tab_name)

    for disc in disciplines:
        disc_name = disc["name"]
        short = get_short_name(disc_name)
        ctype = disc.get("control_type", "credit")
        v_title = f"В_{short}" if ctype != "course_project" else f"В_{short}_КР"
        _create_v_report_sheet(
            worksheet=sheets_map[v_title],
            students=students,
            discipline=disc,
            disciplines=disciplines,
            group_name=group_name,
            first_pc_tab_name=first_pc_tab_name,
        )

    _create_consolidated_dashboard(sheets_map["Відомості"], students, disciplines, group_name, first_pc_tab_name)

    title = f"Журнал_{group_name}"
    safe_filename = re.sub(r'[\\/*?:"<>|]', "_", title) + ".xlsx"

    out_dir = Path(output_dir) if output_dir else Path(tempfile.mkdtemp())
    out_dir.mkdir(parents=True, exist_ok=True)
    file_path = out_dir / safe_filename
    wb.save(str(file_path))

    return {
        "title": title,
        "file_path": str(file_path),
        "filename": safe_filename,
    }
